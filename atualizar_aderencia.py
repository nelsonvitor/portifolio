# -*- coding: utf-8 -*-
"""
atualizar_aderencia.py

Rotina diária do relatório de aderência de produção:
  1. Copia a aba "ADERÊNCIA" de "ADERÊNCIA DE PRODUÇÃO.xlsx" (só os valores, sem fórmulas)
     para a pasta do projeto, com o nome "base_aderencia.xlsx".
  2. Se a base mudou, faz commit e envia ao GitHub (git pull --rebase + git push).

Uso:
    python atualizar_aderencia.py              rotina completa
    python atualizar_aderencia.py --sem-push   só atualiza o arquivo, sem git
    python atualizar_aderencia.py --forcar     ignora a trava de segurança (base muito menor que a atual)

Requisito: o pacote openpyxl (o script instala sozinho na primeira execução, se faltar)
Log: <pasta do projeto>\\logs\\aderencia.log
"""
import argparse
import datetime as dt
import logging
import os
import shutil
import subprocess
import sys
import tempfile
import time
import unicodedata
from logging.handlers import RotatingFileHandler
from pathlib import Path

# ----------------------------------------------------------------------------
# CONFIGURAÇÃO
# ----------------------------------------------------------------------------
ORIGEM_PASTA = r"C:\Users\jaildo.junior\Grupo RFK\Industria - Documents\Novo - 1921680101\PCP\Programação\CONTROLE DE PRODUÇÃO"
ORIGEM_ARQUIVO = "ADERÊNCIA DE PRODUÇÃO.xlsx"
ABA = "ADERÊNCIA"

DESTINO_PASTA = r"C:\Users\jaildo.junior\Desktop\ANALISE_PCP_RFK"   # pasta do repositório git
DESTINO_ARQUIVO = "base_aderencia.xlsx"

# Colunas que o relatório precisa (o cabeçalho é localizado por elas)
COLUNAS_OBRIGATORIAS = ["PRODUTO", "LOCAL", "ESTOQUE INICIAL", "COTA LOCAL", "PRODUÇÃO REAL"]

# Trava: se a nova base tiver menos que essa fração das linhas da base atual, não publica
FRACAO_MINIMA_LINHAS = 0.5

EXTENSAO_LOG = "aderencia.log"
# ----------------------------------------------------------------------------

log = logging.getLogger("aderencia")
NO_WINDOW = 0x08000000 if os.name == "nt" else 0   # evita janelas piscando ao rodar com pythonw


class Erro(Exception):
    """Erro esperado, com mensagem clara para o log."""


def configurar_log():
    pasta_logs = Path(DESTINO_PASTA) / "logs"
    pasta_logs.mkdir(parents=True, exist_ok=True)
    log.setLevel(logging.INFO)
    fmt = logging.Formatter("%(asctime)s  %(levelname)-7s %(message)s", "%d/%m/%Y %H:%M:%S")
    arq = RotatingFileHandler(pasta_logs / EXTENSAO_LOG, maxBytes=1_000_000, backupCount=3, encoding="utf-8")
    arq.setFormatter(fmt)
    log.addHandler(arq)
    if sys.stdout is not None:            # pythonw não tem console
        try:
            sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass
        console = logging.StreamHandler(sys.stdout)
        console.setFormatter(fmt)
        log.addHandler(console)


def norm(v):
    """Compara textos sem acento, sem diferença de maiúsculas e sem espaços sobrando."""
    if v is None:
        return ""
    t = unicodedata.normalize("NFD", str(v))
    t = "".join(c for c in t if unicodedata.category(c) != "Mn")
    return " ".join(t.upper().split())


def garantir_openpyxl():
    """Se o Python que está rodando o script não tem o openpyxl, instala nele mesmo (só na 1ª vez)."""
    try:
        import openpyxl  # noqa: F401
        return
    except ImportError:
        pass
    exe = sys.executable
    if exe.lower().endswith("pythonw.exe"):              # pip funciona melhor com o python.exe irmão
        irmao = exe[: -len("pythonw.exe")] + "python.exe"
        if os.path.exists(irmao):
            exe = irmao
    log.warning("openpyxl não está instalado neste Python (%s). Instalando...", exe)
    try:
        r = subprocess.run([exe, "-m", "pip", "install", "--disable-pip-version-check", "openpyxl"],
                           capture_output=True, text=True, encoding="utf-8", errors="replace",
                           timeout=600, creationflags=NO_WINDOW)
    except Exception as e:
        raise Erro(f"Não consegui executar o pip ({exe}): {e}")
    if r.returncode != 0:
        detalhe = ((r.stdout or "") + (r.stderr or "")).strip()[-1500:]
        raise Erro(f"Falha ao instalar o openpyxl automaticamente (código {r.returncode}):\n{detalhe}")
    import importlib
    importlib.invalidate_caches()
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        raise Erro("openpyxl foi instalado, mas ainda não pode ser importado. Rode a tarefa de novo.")
    log.info("openpyxl instalado com sucesso.")


# ----------------------------------------------------------------------------
# 1) LER A ABA DA PLANILHA DE ORIGEM
# ----------------------------------------------------------------------------
def copiar_origem(origem: Path) -> Path:
    """Copia para uma pasta temporária (evita problemas com arquivo aberto/sincronizando)."""
    tmp_dir = Path(tempfile.mkdtemp(prefix="aderencia_"))
    destino = tmp_dir / "origem.xlsx"
    ultimo = None
    for tentativa in range(1, 4):
        try:
            shutil.copy2(origem, destino)
            return destino
        except (PermissionError, OSError) as e:
            ultimo = e
            log.warning("Não consegui copiar a origem (tentativa %d/3): %s", tentativa, e)
            time.sleep(5)
    raise Erro(f"Não foi possível ler a planilha de origem (aberta/bloqueada?): {ultimo}")


def ler_aba(caminho: Path):
    import openpyxl
    wb = openpyxl.load_workbook(caminho, data_only=True, read_only=True)
    try:
        nome = next((n for n in wb.sheetnames if norm(n) == norm(ABA)), None)
        if nome is None:
            raise Erro(f'A aba "{ABA}" não existe. Abas encontradas: {", ".join(wb.sheetnames)}')
        linhas, vazias = [], 0
        for row in wb[nome].iter_rows(values_only=True):
            row = list(row)
            if all(c is None or (isinstance(c, str) and not c.strip()) for c in row):
                vazias += 1
                if vazias >= 500:      # fim dos dados (evita varrer 1 milhão de linhas vazias)
                    break
                continue
            vazias = 0
            linhas.append(row)
    finally:
        wb.close()
    return linhas


def montar_tabela(linhas):
    """Acha a linha do cabeçalho, remove colunas sem título e linhas vazias."""
    obrig = [norm(c) for c in COLUNAS_OBRIGATORIAS]
    idx = None
    for i, row in enumerate(linhas[:50]):
        cab = {norm(c) for c in row}
        if all(o in cab for o in obrig):
            idx = i
            break
    if idx is None:
        raise Erro("Não encontrei o cabeçalho na aba. Colunas esperadas: " + ", ".join(COLUNAS_OBRIGATORIAS))

    cab = linhas[idx]
    manter = [j for j, c in enumerate(cab) if c is not None and str(c).strip() != ""]
    tabela = [[str(cab[j]).strip() for j in manter]]
    for row in linhas[idx + 1:]:
        vals = [row[j] if j < len(row) else None for j in manter]
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in vals):
            continue
        tabela.append(vals)
    if len(tabela) < 2:
        raise Erro("A aba não tem nenhuma linha de dados abaixo do cabeçalho.")
    return tabela


# ----------------------------------------------------------------------------
# 2) SALVAR base_aderencia.xlsx
# ----------------------------------------------------------------------------
def ler_base_existente(caminho: Path):
    if not caminho.exists():
        return None
    import openpyxl
    try:
        wb = openpyxl.load_workbook(caminho, data_only=True, read_only=True)
        try:
            return [list(r) for r in wb[wb.sheetnames[0]].iter_rows(values_only=True)]
        finally:
            wb.close()
    except Exception as e:                       # base antiga corrompida: trata como inexistente
        log.warning("Não consegui ler a base atual (%s); ela será substituída.", e)
        return None


def salvar_base(tabela, destino: Path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = ABA
    for row in tabela:
        ws.append(row)
    tmp = destino.with_name(destino.stem + ".tmp.xlsx")
    wb.save(tmp)
    try:
        os.replace(tmp, destino)                 # troca atômica
    except PermissionError:
        tmp.unlink(missing_ok=True)
        raise Erro(f"Não consegui gravar {destino.name}: feche o arquivo se ele estiver aberto no Excel.")


# ----------------------------------------------------------------------------
# 3) GIT
# ----------------------------------------------------------------------------
def git(*args, ok_codes=(0,)):
    env = dict(os.environ, GIT_TERMINAL_PROMPT="0")     # nunca ficar esperando senha
    r = subprocess.run(["git", *args], cwd=DESTINO_PASTA, capture_output=True, text=True,
                       encoding="utf-8", errors="replace", env=env, timeout=180,
                       creationflags=NO_WINDOW)
    saida = (r.stdout or "").strip() + ("\n" + r.stderr.strip() if r.stderr and r.stderr.strip() else "")
    if r.returncode not in ok_codes:
        raise Erro(f'git {" ".join(args)} falhou (código {r.returncode}):\n{saida}')
    return (r.stdout or "").strip()


def publicar(alterou: bool):
    try:
        git("rev-parse", "--is-inside-work-tree")
    except (Erro, FileNotFoundError) as e:
        if isinstance(e, FileNotFoundError):
            raise Erro("Git não encontrado. Instale o Git for Windows e reinicie o computador.")
        raise Erro(f"{DESTINO_PASTA} não é um repositório git.")

    branch = git("rev-parse", "--abbrev-ref", "HEAD")
    if branch == "HEAD":
        raise Erro("O repositório está em 'detached HEAD'. Volte para uma branch (ex.: git checkout main).")

    if alterou:
        git("add", "--", DESTINO_ARQUIVO)
        if git("status", "--porcelain", "--", DESTINO_ARQUIVO):
            msg = f"Atualiza base_aderencia ({dt.datetime.now():%d/%m/%Y %H:%M})"
            git("commit", "-m", msg, "--", DESTINO_ARQUIVO)      # só este arquivo entra no commit
            log.info("Commit criado: %s", msg)

    try:
        git("pull", "--rebase", "--autostash", "origin", branch)
    except Erro:
        try:
            git("rebase", "--abort", ok_codes=(0, 128))
        except Erro:
            pass
        raise
    git("push", "origin", branch)
    log.info("Push concluído (branch %s).", branch)


# ----------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser(description="Atualiza base_aderencia.xlsx e publica no GitHub.")
    ap.add_argument("--sem-push", action="store_true", help="não executa git")
    ap.add_argument("--forcar", action="store_true", help="ignora a trava de quantidade de linhas")
    args = ap.parse_args()

    configurar_log()
    log.info("=== Início da atualização ===")
    try:
        garantir_openpyxl()
        origem = Path(ORIGEM_PASTA) / ORIGEM_ARQUIVO
        if not origem.exists():
            raise Erro(f"Planilha de origem não encontrada: {origem}")
        mod = dt.datetime.fromtimestamp(origem.stat().st_mtime)
        log.info("Origem: %s (salva em %s)", origem.name, mod.strftime("%d/%m/%Y %H:%M"))

        copia = copiar_origem(origem)
        try:
            tabela = montar_tabela(ler_aba(copia))
        finally:
            shutil.rmtree(copia.parent, ignore_errors=True)
        n_novo = len(tabela) - 1
        log.info("Aba %s lida: %d linhas de dados, %d colunas.", ABA, n_novo, len(tabela[0]))

        destino = Path(DESTINO_PASTA) / DESTINO_ARQUIVO
        atual = ler_base_existente(destino)
        alterou = True
        if atual is not None:
            n_atual = max(len(atual) - 1, 0)
            if not args.forcar and n_atual > 0 and n_novo < n_atual * FRACAO_MINIMA_LINHAS:
                raise Erro(f"Trava de segurança: a nova base tem {n_novo} linhas e a atual tem {n_atual}. "
                           f"Nada foi alterado. Se for isso mesmo, rode com --forcar.")
            alterou = [list(r) for r in tabela] != [list(r) for r in atual]

        if alterou:
            salvar_base(tabela, destino)
            log.info("%s atualizado.", DESTINO_ARQUIVO)
        else:
            log.info("Sem mudança nos dados desde a última execução; arquivo mantido.")

        if args.sem_push:
            log.info("--sem-push: git ignorado.")
        else:
            publicar(alterou)
        log.info("=== Fim: OK ===")
        return 0
    except Erro as e:
        log.error("%s", e)
    except Exception:
        log.exception("Erro inesperado")
    log.error("=== Fim: COM ERRO ===")
    return 1


if __name__ == "__main__":
    sys.exit(main())
